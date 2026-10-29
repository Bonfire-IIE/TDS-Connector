from __future__ import annotations

import hashlib
import json
import math
import tempfile
from dataclasses import replace
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import duckdb
import openpyxl
import pyarrow.orc as arrow_orc
import pyarrow.parquet as arrow_parquet

from ..domain.catalog import ColumnSchema, InspectionResult
from ..domain.errors import DomainError


def inspect_path(path: Path, object_ref: str, sample_rows: int, options: dict[str, Any] | None = None) -> InspectionResult:
    options = options or {}
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        merged = {**options}
        if suffix == ".tsv" and merged.get("delimiter") in {None, "", "auto"}:
            merged["delimiter"] = "\t"
        return _duckdb_file(path, object_ref, "tsv" if suffix == ".tsv" else "csv", sample_rows, merged)
    if suffix in {".json", ".jsonl", ".ndjson"}:
        return _duckdb_file(path, object_ref, "jsonl" if suffix != ".json" else "json", sample_rows, options)
    if suffix == ".parquet":
        return _arrow_file(path, object_ref, "parquet", sample_rows)
    if suffix == ".orc":
        return _arrow_file(path, object_ref, "orc", sample_rows)
    if suffix in {".xlsx", ".xlsm"}:
        return _excel(path, object_ref, sample_rows, options)
    raise DomainError(422, f"暂不支持探测该文件格式: {suffix or 'unknown'}")


def workbook_sheets(path: Path) -> list[str]:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    except Exception as exc:
        raise DomainError(422, f"无法读取 Excel 工作簿: {exc}") from exc


def _duckdb_file(path: Path, object_ref: str, fmt: str, sample_rows: int,
                 options: dict[str, Any]) -> InspectionResult:
    connection = duckdb.connect(":memory:")
    temporary_path: Path | None = None
    try:
        if fmt in {"csv", "tsv"}:
            input_path, detected_encoding, temporary_path = _normalized_delimited_file(path, options)
            header = options.get("header", "auto")
            kwargs: dict[str, Any] = {"sample_size": -1}
            if header in {True, False}:
                kwargs["header"] = header
            if options.get("delimiter") not in {None, "", "auto"}:
                kwargs["delimiter"] = str(options["delimiter"])
            relation = connection.read_csv(str(input_path), **kwargs)
        else:
            detected_encoding = None
            relation = connection.read_json(str(path), format="newline_delimited" if fmt == "jsonl" else "auto")
        columns = [ColumnSchema(name=name, type=_duck_type(str(dtype)), physical_type=str(dtype),
                                nullable=True, confidence=0.9) for name, dtype in zip(relation.columns, relation.types)]
        rows = [[_json_value(value) for value in row] for row in relation.limit(sample_rows).fetchall()]
        result = _result(object_ref, fmt, "inferred", columns, rows, len(rows) < sample_rows, path)
        if detected_encoding and detected_encoding != "utf-8":
            return replace(result, warnings=[f"检测到源文件编码为 {detected_encoding}，结构探测时已临时转换为 UTF-8"])
        return result
    except DomainError:
        raise
    except Exception as exc:
        raise DomainError(422, f"结构探测失败: {str(exc).splitlines()[0]}") from exc
    finally:
        connection.close()
        if temporary_path:
            temporary_path.unlink(missing_ok=True)


def _normalized_delimited_file(path: Path, options: dict[str, Any]) -> tuple[Path, str, Path | None]:
    """Return a UTF-8 view of a delimited text file without changing the source."""
    requested = str(options.get("encoding") or "").strip().lower().replace("_", "-")
    aliases = {"utf8": "utf-8", "utf-8-sig": "utf-8-sig", "gbk": "gb18030", "gb2312": "gb18030"}
    requested = aliases.get(requested, requested)
    candidates = [requested] if requested else ["utf-8-sig", "gb18030"]
    raw = path.read_bytes()
    decoded: str | None = None
    detected = ""
    errors: list[str] = []
    for encoding in candidates:
        if encoding not in {"utf-8", "utf-8-sig", "gb18030"}:
            raise DomainError(422, f"暂不支持 CSV/TSV 字符编码: {encoding}")
        try:
            decoded = raw.decode(encoding)
            detected = "utf-8" if encoding in {"utf-8", "utf-8-sig"} else encoding
            break
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: 第 {exc.start + 1} 字节")
    if decoded is None:
        detail = "；".join(errors)
        raise DomainError(422, f"无法识别 CSV/TSV 字符编码（{detail}）")
    if detected == "utf-8" and not raw.startswith(b"\xef\xbb\xbf"):
        return path, detected, None
    handle = tempfile.NamedTemporaryFile(prefix="bonfire-inspect-text-", suffix=path.suffix, delete=False)
    temporary = Path(handle.name)
    try:
        with handle:
            handle.write(decoded.encode("utf-8"))
        return temporary, detected, temporary
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _arrow_file(path: Path, object_ref: str, fmt: str, sample_rows: int) -> InspectionResult:
    try:
        table = arrow_parquet.read_table(path) if fmt == "parquet" else arrow_orc.ORCFile(path).read()
        columns = [ColumnSchema(field.name, _arrow_type(str(field.type)), str(field.type), field.nullable, 1.0)
                   for field in table.schema]
        rows = [[_json_value(value) for value in row.values()]
                for row in table.slice(0, sample_rows).to_pylist()]
        return _result(object_ref, fmt, "declared", columns, rows, table.num_rows <= sample_rows, path)
    except Exception as exc:
        raise DomainError(422, f"{fmt.upper()} 结构探测失败: {str(exc).splitlines()[0]}") from exc


def _excel(path: Path, object_ref: str, sample_rows: int, options: dict[str, Any]) -> InspectionResult:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = str(options.get("sheet") or "")
        if not sheet_name:
            if len(workbook.sheetnames) != 1:
                raise DomainError(422, "Excel 工作簿包含多个工作表，请先选择 sheet")
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise DomainError(404, "Excel 工作表不存在")
        rows = list(workbook[sheet_name].iter_rows(values_only=True, max_row=sample_rows + 1))
        if not rows:
            raise DomainError(422, "Excel 工作表为空")
        header = _headers(rows[0])
        data = rows[1:]
        columns = []
        for index, name in enumerate(header):
            values = [row[index] for row in data if index < len(row) and row[index] is not None]
            columns.append(ColumnSchema(name, _python_type(values), None, len(values) < len(data), 0.8))
        sample = [[_json_value(value) for value in row] for row in data]
        result = _result(object_ref, "xlsx", "inferred", columns, sample, len(rows) <= sample_rows, path)
        return InspectionResult(**{**result.__dict__, "warnings": [f"sheet={sheet_name}"]})
    finally:
        workbook.close()


def _headers(values: tuple[Any, ...]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for index, value in enumerate(values, 1):
        name = str(value).strip() if value is not None else ""
        name = name or f"column_{index}"
        base = name
        suffix = 2
        while name in seen:
            name = f"{base}_{suffix}"
            suffix += 1
        seen.add(name)
        result.append(name)
    return result


def _python_type(values: list[Any]) -> str:
    kinds = {type(value) for value in values}
    if not kinds:
        return "string"
    if kinds <= {bool}:
        return "boolean"
    if kinds <= {int, bool}:
        return "bigint"
    if kinds <= {int, float, bool}:
        return "double"
    if kinds <= {datetime}:
        return "timestamp"
    if kinds <= {date}:
        return "date"
    if kinds <= {time}:
        return "time"
    return "string"


def _duck_type(value: str) -> str:
    upper = value.upper()
    if upper == "BOOLEAN": return "boolean"
    if upper in {"TINYINT", "UTINYINT"}: return "tinyint"
    if upper in {"SMALLINT", "USMALLINT"}: return "smallint"
    if upper in {"INTEGER", "UINTEGER"}: return "int"
    if upper in {"BIGINT", "UBIGINT", "HUGEINT"}: return "bigint"
    if upper in {"FLOAT", "REAL"}: return "float"
    if upper == "DOUBLE": return "double"
    if upper.startswith("DECIMAL"):
        return upper.lower()
    if upper == "DATE": return "date"
    if upper == "TIME": return "time"
    if "TIMESTAMP" in upper: return "timestamp"
    if upper in {"BLOB", "BYTEA"}: return "bytes"
    return "string"


def _arrow_type(value: str) -> str:
    lower = value.lower()
    if lower == "bool": return "boolean"
    if lower in {"int8", "uint8"}: return "tinyint"
    if lower in {"int16", "uint16"}: return "smallint"
    if lower in {"int32", "uint32"}: return "int"
    if lower in {"int64", "uint64"}: return "bigint"
    if lower == "float": return "float"
    if lower == "double": return "double"
    if lower.startswith("decimal"):
        numbers = [int(item) for item in __import__("re").findall(r"\d+", lower)]
        return f"decimal({numbers[0]},{numbers[1]})" if len(numbers) >= 2 and numbers[0] <= 38 else "string"
    if lower.startswith("date"): return "date"
    if lower.startswith("time") and not lower.startswith("timestamp"): return "time"
    if lower.startswith("timestamp"): return "timestamp"
    if lower in {"binary", "large_binary", "fixed_size_binary"}: return "bytes"
    return "string"


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return f"<bytes:{len(value)}>"
    return str(value)


def _result(object_ref: str, fmt: str, source: str, columns: list[ColumnSchema], rows: list[list[Any]],
            complete: bool, path: Path) -> InspectionResult:
    stat = path.stat()
    signature = json.dumps({"size": stat.st_size, "mtime_ns": stat.st_mtime_ns,
                            "columns": [column.__dict__ for column in columns]}, sort_keys=True).encode()
    return InspectionResult(object_ref, "structured", fmt, source, columns, len(rows), complete, rows,
                            "sha256:" + hashlib.sha256(signature).hexdigest())
